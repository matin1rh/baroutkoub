from typing import List, Dict, Any, Tuple
import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Font, Alignment
from src.normalization import (
    parse_amount,
    parse_jalali_date,
    normalize_digits,
    normalize_persian_text
)

GREEN_FILL = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
RED_FILL = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
HEADER_FILL = PatternFill(start_color="D6D8DB", end_color="D6D8DB", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style='thin', color='C0C0C0'),
    right=Side(style='thin', color='C0C0C0'),
    top=Side(style='thin', color='C0C0C0'),
    bottom=Side(style='thin', color='C0C0C0')
)

class ExcelProcessor:
    @staticmethod
    def _find_header_row(rows: List[Tuple], keywords: List[str]) -> Tuple[int, List[str]]:
        for r_idx, row in enumerate(rows[:10]):
            row_str = " ".join([normalize_persian_text(str(c)) for c in row if c is not None])
            match_count = sum(1 for kw in keywords if normalize_persian_text(kw) in row_str)
            if match_count >= 2:
                return r_idx, [normalize_persian_text(str(c)) if c is not None else "" for c in row]
        return 0, [normalize_persian_text(str(c)) if c is not None else "" for c in rows[0]]

    @staticmethod
    def read_system_file(file_path: str) -> List[Dict[str, Any]]:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        header_idx, header = ExcelProcessor._find_header_row(
            rows, ["بستانکار", "بدهکار", "نام حساب", "تاریخ", "کد رهگیری", "سند"]
        )

        col_keywords = {
            "credit": ["بستانکار مالی", "بستانکار"],
            "debit": ["بدهکار مالی", "بدهکار"],
            "op_date": ["تاریخ عملیات", "تاريخ عمليات"],
            "doc_date": ["تاریخ سند", "تاريخ سند"],
            "check_date": ["تاریخ چک", "تاريخ چک"],
            "account_name": ["نام حساب", "طرف حساب", "حساب"],
            "doc_type": ["نوع سند", "شرح سند"],
            "tracking": ["شماره چک / کد رهگیری", "کد رهگیری", "شماره چک", "رهگیری", "شماره سند"]
        }

        idx_map = {}
        for key, kws in col_keywords.items():
            for kw in kws:
                n_kw = normalize_persian_text(kw)
                for i, h in enumerate(header):
                    if n_kw in h:
                        idx_map[key] = i
                        break
                if key in idx_map:
                    break

        records = []
        for r_idx, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
            if not any(row):
                continue

            credit = parse_amount(row[idx_map["credit"]]) if "credit" in idx_map and idx_map["credit"] < len(row) else 0
            debit = parse_amount(row[idx_map["debit"]]) if "debit" in idx_map and idx_map["debit"] < len(row) else 0

            if credit:
                amount = credit
                direction = "CREDIT"
            elif debit:
                amount = debit
                direction = "DEBIT"
            else:
                continue

            raw_date = None
            for d_col in ("op_date", "doc_date", "check_date"):
                if d_col in idx_map and idx_map[d_col] < len(row) and row[idx_map[d_col]]:
                    raw_date = row[idx_map[d_col]]
                    break

            j_date = parse_jalali_date(raw_date)
            tracking = normalize_digits(str(row[idx_map["tracking"]] or "")) if "tracking" in idx_map and idx_map["tracking"] < len(row) else ""
            acc_name = str(row[idx_map["account_name"]] or "").strip() if "account_name" in idx_map and idx_map["account_name"] < len(row) else ""
            doc_type = str(row[idx_map["doc_type"]] or "").strip() if "doc_type" in idx_map and idx_map["doc_type"] < len(row) else ""

            records.append({
                "sys_index": len(records),
                "original_row": r_idx,
                "amount": amount,
                "direction": direction,
                "date": j_date,
                "tracking_code": tracking,
                "account_name": acc_name,
                "doc_type": doc_type,
                "raw_desc": f"{acc_name} {doc_type} {tracking}".strip(),
                "raw_data": list(row)
            })

        return records

    @staticmethod
    def read_bank_file(file_path: str) -> List[Dict[str, Any]]:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        header_idx, header = ExcelProcessor._find_header_row(
            rows, ["مبلغ گردش بستانکار", "مبلغ گردش بدهکار", "شرح", "واریز کننده", "سریال", "شناسه واریز"]
        )

        col_keywords = {
            "credit": ["مبلغ گردش بستانکار", "گردش بستانکار", "بستانکار"],
            "debit": ["مبلغ گردش بدهکار", "گردش بدهکار", "بدهکار"],
            "desc": ["شرح", "شرح تراکنش"],
            "party": ["واریز کننده/ ذیتفع", "واریز کننده", "ذینفع", "طرف تراکنش"],
            "serial": ["شماره سریال", "سریال", "شماره پیگیری"],
            "deposit_id": ["شناسه واریز", "شناسه"],
            "date": ["تاریخ", "تاریخ تراکنش"]
        }

        idx_map = {}
        for key, kws in col_keywords.items():
            for kw in kws:
                n_kw = normalize_persian_text(kw)
                for i, h in enumerate(header):
                    if n_kw in h:
                        idx_map[key] = i
                        break
                if key in idx_map:
                    break

        records = []
        for r_idx, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
            if not any(row):
                continue

            credit = parse_amount(row[idx_map["credit"]]) if "credit" in idx_map and idx_map["credit"] < len(row) else 0
            debit = parse_amount(row[idx_map["debit"]]) if "debit" in idx_map and idx_map["debit"] < len(row) else 0

            if credit:
                amount = credit
                direction = "CREDIT"
            elif debit:
                amount = debit
                direction = "DEBIT"
            else:
                continue

            raw_date = row[idx_map.get("date", 0)] if "date" in idx_map and idx_map["date"] < len(row) else None
            j_date = parse_jalali_date(raw_date)

            desc = str(row[idx_map["desc"]] or "").strip() if "desc" in idx_map and idx_map["desc"] < len(row) else ""
            party = str(row[idx_map["party"]] or "").strip() if "party" in idx_map and idx_map["party"] < len(row) else ""
            serial = normalize_digits(str(row[idx_map["serial"]] or "")).strip() if "serial" in idx_map and idx_map["serial"] < len(row) else ""
            dep_id = normalize_digits(str(row[idx_map["deposit_id"]] or "")).strip() if "deposit_id" in idx_map and idx_map["deposit_id"] < len(row) else ""

            records.append({
                "bank_index": len(records),
                "original_row": r_idx,
                "amount": amount,
                "direction": direction,
                "date": j_date,
                "description": desc,
                "party_name": party,
                "serial_no": serial,
                "deposit_id": dep_id,
                "raw_desc": f"{party} {desc}".strip(),
                "raw_data": list(row)
            })

        return records

    @staticmethod
    def write_reconciliation_results(
        sys_file: str, bank_file: str,
        sys_records: List[Dict[str, Any]], bank_records: List[Dict[str, Any]],
        match_results: Dict[str, Any],
        out_sys_path: str, out_bank_path: str, out_report_path: str
    ):
        sys_matches = match_results.get("sys_matches", {})
        bank_matches = match_results.get("bank_matches", {})

        # ۱. خروجی مرتب‌شده بر اساس مبلغ برای فایل سیستم
        wb_sys_orig = openpyxl.load_workbook(sys_file, data_only=True)
        ws_orig = wb_sys_orig.active
        orig_headers = [c for c in next(ws_orig.iter_rows(values_only=True))]

        wb_sys = openpyxl.Workbook()
        ws_sys = wb_sys.active
        ws_sys.title = "تطبیق سیستم"
        ws_sys.views.sheetView[0].rightToLeft = True

        headers_sys = list(orig_headers) + ["وضعیت تطبیق", "دلیل تطبیق", "ردیف متناظر بانک"]
        ws_sys.append(headers_sys)
        for col_num in range(1, len(headers_sys) + 1):
            ws_sys.cell(row=1, column=col_num).fill = HEADER_FILL

        # سورت سیستم بر اساس مبلغ (نزولی)
        sorted_sys = sorted(sys_records, key=lambda x: x["amount"], reverse=True)

        for s_rec in sorted_sys:
            s_idx = s_rec["sys_index"]
            match = sys_matches.get(s_idx)
            row_data = list(s_rec["raw_data"])

            if match:
                fill_to_apply = GREEN_FILL if match["status"] == "GREEN" else YELLOW_FILL
                status_text = "تطبیق قطعی" if match["status"] == "GREEN" else "تطبیق احتمالی (بازبینی)"
                reason_text = match["reason"]
                b_idx = match["matched_index"]
                matched_row_str = str(bank_records[b_idx]["original_row"])
            else:
                fill_to_apply = RED_FILL
                status_text = "عدم تطبیق"
                reason_text = "فاقد رکورد متناظر در بانک"
                matched_row_str = "-"

            row_data.extend([status_text, reason_text, matched_row_str])
            ws_sys.append(row_data)

            curr_row = ws_sys.max_row
            for c_idx in range(1, len(row_data) + 1):
                cell = ws_sys.cell(row=curr_row, column=c_idx)
                cell.fill = fill_to_apply
                cell.border = THIN_BORDER

        wb_sys.save(out_sys_path)

        # ۲. خروجی مرتب‌شده بر اساس مبلغ برای فایل بانک
        wb_bank_orig = openpyxl.load_workbook(bank_file, data_only=True)
        ws_b_orig = wb_bank_orig.active
        b_orig_headers = [c for c in next(ws_b_orig.iter_rows(values_only=True))]

        wb_bank = openpyxl.Workbook()
        ws_bank = wb_bank.active
        ws_bank.title = "تطبیق بانک"
        ws_bank.views.sheetView[0].rightToLeft = True

        headers_bank = list(b_orig_headers) + ["وضعیت تطبیق", "دلیل تطبیق", "ردیف متناظر سیست"]
        ws_bank.append(headers_bank)
        for col_num in range(1, len(headers_bank) + 1):
            ws_bank.cell(row=1, column=col_num).fill = HEADER_FILL

        # سورت بانک بر اساس مبلغ (نزولی)
        sorted_bank = sorted(bank_records, key=lambda x: x["amount"], reverse=True)

        for b_rec in sorted_bank:
            b_idx = b_rec["bank_index"]
            match = bank_matches.get(b_idx)
            row_data = list(b_rec["raw_data"])

            if match:
                fill_to_apply = GREEN_FILL if match["status"] == "GREEN" else YELLOW_FILL
                status_text = "تطبیق قطعی" if match["status"] == "GREEN" else "تطبیق احتمالی (بازبینی)"
                reason_text = match["reason"]
                s_idx = match["matched_index"]
                matched_row_str = str(sys_records[s_idx]["original_row"])
            else:
                fill_to_apply = RED_FILL
                status_text = "عدم تطبیق"
                reason_text = "فاقد رکورد متناظر در سیستم"
                matched_row_str = "-"

            row_data.extend([status_text, reason_text, matched_row_str])
            ws_bank.append(row_data)

            curr_row = ws_bank.max_row
            for c_idx in range(1, len(row_data) + 1):
                cell = ws_bank.cell(row=curr_row, column=c_idx)
                cell.fill = fill_to_apply
                cell.border = THIN_BORDER

        wb_bank.save(out_bank_path)

        # ۳. گزارش خلاصه مرتب‌شده
        wb_rep = openpyxl.Workbook()
        ws_rep = wb_rep.active
        ws_rep.title = "خلاصه مغایرت‌گیری"
        ws_rep.views.sheetView[0].rightToLeft = True

        green_count = sum(1 for m in sys_matches.values() if m.get("status") == "GREEN")
        yellow_count = sum(1 for m in sys_matches.values() if m.get("status") == "YELLOW")
        red_sys_count = len(sys_records) - len(sys_matches)

        summary_rows = [
            ["شاخص آماری", "مقدار"],
            ["کل تراکنش‌های سیستم", len(sys_records)],
            ["کل تراکنش‌های بانک", len(bank_records)],
            ["تطبیق‌های قطعی (سبز)", green_count],
            ["تطبیق‌های نیازمند بررسی (زرد)", yellow_count],
            ["اقلام باز و بدون تطبیق سیستم (قرمز)", red_sys_count],
            ["درصد موفقیت تطبیق", f"{round(((green_count + yellow_count) / max(len(sys_records), 1)) * 100, 2)}%"]
        ]
        for row_data in summary_rows:
            ws_rep.append(row_data)

        # برگه جزئیات تراکنش‌ها مرتب‌شده بر مبنای مبلغ
        ws_det = wb_rep.create_sheet(title="جزئیات تمام تطبیق‌ها")
        ws_det.views.sheetView[0].rightToLeft = True
        ws_det.append(["ردیف سیست", "ردیف بانک", "مبلغ", "تاریخ سیست", "تاریخ بانک", "درصد اطمینان", "وضعیت", "دلیل"])
        
        matched_pairs = []
        for s_idx, match in sys_matches.items():
            b_idx = match["matched_index"]
            s_rec = sys_records[s_idx]
            b_rec = bank_records[b_idx]
            matched_pairs.append((s_rec, b_rec, match))

        # سورت جزئیات گزارش بر اساس مبلغ
        matched_pairs.sort(key=lambda item: item[0]["amount"], reverse=True)

        for s_rec, b_rec, match in matched_pairs:
            ws_det.append([
                s_rec["original_row"],
                b_rec["original_row"],
                f"{s_rec['amount']:,}",
                str(s_rec["date"] or "-"),
                str(b_rec["date"] or "-"),
                f"{int(match.get('confidence', 1.0) * 100)}%",
                "سبز" if match.get("status") == "GREEN" else "زرد",
                match.get("reason", "")
            ])

        wb_rep.save(out_report_path)